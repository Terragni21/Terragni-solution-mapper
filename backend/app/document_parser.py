"""Build a searchable index from the Terragni source library.

Supported source types:
- PDF  -> page-aware text
- PPTX -> slide-aware text
- DOCX -> heading/paragraph/table text
- MD/TXT -> text

The originals are never changed. The index is rebuilt automatically whenever a
source file is added, removed or changed.
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Iterable

from pypdf import PdfReader
from pptx import Presentation
from docx import Document

from .config import INDEX_PATH, RAW_SOURCES_DIR, CASE_STUDIES_DIR, MAX_CHUNK_CHARS

SUPPORTED = {".pdf", ".pptx", ".docx", ".md", ".txt"}


def source_paths() -> list[Path]:
    paths: list[Path] = []
    for directory in (RAW_SOURCES_DIR, CASE_STUDIES_DIR):
        if not directory.exists():
            continue
        paths.extend(
            p for p in directory.rglob("*")
            if p.is_file() and p.suffix.lower() in SUPPORTED
        )
    return sorted(paths, key=lambda p: str(p).lower())


def fingerprint(path: Path) -> str:
    h = hashlib.sha256()
    h.update(str(path.relative_to(path.parents[2] if len(path.parents) > 2 else path.parent)).encode())
    h.update(str(path.stat().st_size).encode())
    h.update(str(path.stat().st_mtime_ns).encode())
    return h.hexdigest()


def clean(text: str) -> str:
    text = (text or "").replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str, max_chars: int = MAX_CHUNK_CHARS) -> list[str]:
    text = clean(text)
    if not text:
        return []
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks, current = [], ""
    for para in paragraphs:
        if len(current) + len(para) + 2 <= max_chars:
            current = f"{current}\n\n{para}".strip()
        else:
            if current:
                chunks.append(current)
            while len(para) > max_chars:
                cut = para.rfind(" ", 0, max_chars)
                cut = cut if cut > max_chars // 2 else max_chars
                chunks.append(para[:cut].strip())
                para = para[cut:].strip()
            current = para
    if current:
        chunks.append(current)
    return chunks


def parse_pdf(path: Path) -> Iterable[tuple[str, str]]:
    reader = PdfReader(str(path))
    for page_no, page in enumerate(reader.pages, 1):
        text = clean(page.extract_text() or "")
        if text:
            yield f"page {page_no}", text


def parse_pptx(path: Path) -> Iterable[tuple[str, str]]:
    prs = Presentation(str(path))
    for slide_no, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        text = clean("\n".join(parts))
        if text:
            yield f"slide {slide_no}", text


def parse_docx(path: Path) -> Iterable[tuple[str, str]]:
    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = clean(paragraph.text)
        if text:
            parts.append(text)
    for table_no, table in enumerate(doc.tables, 1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(clean(cell.text) for cell in row.cells))
        table_text = clean("\n".join(rows))
        if table_text:
            parts.append(f"Table {table_no}\n{table_text}")
    if parts:
        # Keep DOCX locators simple and stable. Chunks remain traceable to the document.
        yield "document", "\n\n".join(parts)


def parse_text(path: Path) -> Iterable[tuple[str, str]]:
    yield "document", path.read_text(encoding="utf-8", errors="ignore")


def parse_file(path: Path) -> list[dict]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        sections = parse_pdf(path)
    elif ext == ".pptx":
        sections = parse_pptx(path)
    elif ext == ".docx":
        sections = parse_docx(path)
    else:
        sections = parse_text(path)

    records = []
    for locator, text in sections:
        for i, chunk in enumerate(chunk_text(text)):
            records.append({
                "id": f"{path.name}:{locator}:{i + 1}",
                "source": path.name,
                "type": ext[1:],
                "locator": locator,
                "text": chunk,
            })
    return records


def build_index() -> dict:
    documents = []
    fingerprints = {}
    for path in source_paths():
        documents.extend(parse_file(path))
        fingerprints[str(path)] = fingerprint(path)

    index = {
        "version": 2,
        "documents": documents,
        "fingerprints": fingerprints,
    }
    INDEX_PATH.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    return index


def load_index() -> dict:
    current = {str(path): fingerprint(path) for path in source_paths()}

    if not INDEX_PATH.exists():
        return build_index()

    try:
        index = json.loads(INDEX_PATH.read_text(encoding="utf-8"))
        if index.get("version") != 2 or index.get("fingerprints") != current:
            return build_index()
        return index
    except Exception:
        return build_index()


if __name__ == "__main__":
    index = build_index()
    print(f"Indexed {len(index['documents'])} chunks from {len(index['fingerprints'])} source files.")

