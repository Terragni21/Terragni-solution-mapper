from __future__ import annotations

from pathlib import Path

from docx import Document


def extract_docx(path: str | Path) -> list[dict]:
    """
    Extract a DOCX document while preserving
    paragraph and table content.

    Returns:
        [
            {
                "section_order": 1,
                "section_title": "",
                "text": "...",
            },
            ...
        ]
    """

    docx_path = Path(path)

    if not docx_path.exists():
        raise FileNotFoundError(
            f"DOCX file not found: {docx_path}"
        )

    if docx_path.suffix.lower() != ".docx":
        raise ValueError(
            f"Expected a DOCX file, got: {docx_path.suffix}"
        )

    document = Document(
        str(docx_path)
    )

    sections: list[dict] = []

    section_order = 0
    current_title = ""

    # -----------------------------
    # Paragraphs
    # -----------------------------
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if not text:
            continue

        style_name = (
            paragraph.style.name
            if paragraph.style
            else ""
        )

        if style_name.startswith("Heading"):
            current_title = text
            continue

        section_order += 1

        sections.append(
            {
                "section_order": section_order,
                "section_title": current_title,
                "text": text,
            }
        )

    # -----------------------------
    # Tables
    # -----------------------------
    for table in document.tables:

        for row in table.rows:

            values = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    values.append(cell_text)

            if not values:
                continue

            section_order += 1

            sections.append(
                {
                    "section_order": section_order,
                    "section_title": current_title,
                    "text": " | ".join(values),
                }
            )

    return sections