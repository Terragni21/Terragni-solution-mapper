from __future__ import annotations

from pathlib import Path

from pptx import Presentation


def extract_pptx(path: str | Path) -> list[dict]:
    """
    Extract PowerPoint content slide-by-slide.

    Extracts:
    - text boxes
    - titles
    - table cell text

    Returns:
        [
            {
                "slide_number": 1,
                "text": "...",
            },
            ...
        ]
    """

    pptx_path = Path(path)

    if not pptx_path.exists():
        raise FileNotFoundError(
            f"PPTX file not found: {pptx_path}"
        )

    if pptx_path.suffix.lower() != ".pptx":
        raise ValueError(
            f"Expected a PPTX file, got: {pptx_path.suffix}"
        )

    presentation = Presentation(
        str(pptx_path)
    )

    sections: list[dict] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        text_parts: list[str] = []

        for shape in slide.shapes:

            # -----------------------------
            # Normal text-containing shapes
            # -----------------------------
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    text_parts.append(text)

            # -----------------------------
            # Tables
            # -----------------------------
            if getattr(shape, "has_table", False):

                for row in shape.table.rows:
                    row_values = []

                    for cell in row.cells:
                        cell_text = cell.text.strip()

                        if cell_text:
                            row_values.append(
                                cell_text
                            )

                    if row_values:
                        text_parts.append(
                            " | ".join(row_values)
                        )

        slide_text = "\n".join(
            text_parts
        ).strip()

        sections.append(
            {
                "slide_number": slide_number,
                "text": slide_text,
            }
        )

    return sections